"""
Generate a .docx file from a markdown source.

Primary path: uses python-docx if installed.
Fallback: pure-stdlib WordprocessingML writer (minimal) so script still works
even when python-docx is unavailable in sandboxed environments.
"""

import sys
import re
import pathlib
import zipfile
import xml.etree.ElementTree as ET

try:
    from docx import Document  # type: ignore
except Exception:  # pragma: no cover - fallback to stdlib
    Document = None


def parse_markdown(md_text: str):
    """Very small markdown parser for headings, bullets, tables, paragraphs."""
    elements = []
    lines = md_text.splitlines()
    i = 0
    buffer = []

    def flush_paragraph():
        if buffer:
            text = " ".join(buffer).strip()
            if text:
                elements.append({"type": "paragraph", "text": text})
            buffer.clear()

    while i < len(lines):
        line = lines[i].rstrip()

        # Table detection (header + separator)
        if "|" in line and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?\s*$", lines[i + 1]):
            flush_paragraph()
            header = [c.strip() for c in line.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and "|" in lines[i]:
                row = [c.strip() for c in lines[i].strip("|").split("|")]
                rows.append(row)
                i += 1
            elements.append({"type": "table", "header": header, "rows": rows})
            continue

        # Heading
        m = re.match(r"^(#+)\s+(.*)$", line)
        if m:
            flush_paragraph()
            level = min(len(m.group(1)), 6)
            elements.append({"type": "heading", "level": level, "text": m.group(2).strip()})
            i += 1
            continue

        # Bullet
        if line.lstrip().startswith("- "):
            flush_paragraph()
            items = []
            while i < len(lines) and lines[i].lstrip().startswith("- "):
                items.append(lines[i].lstrip()[2:].strip())
                i += 1
            elements.append({"type": "bullet", "items": items})
            continue

        # Blank line
        if not line.strip():
            flush_paragraph()
            i += 1
            continue

        buffer.append(line)
        i += 1

    flush_paragraph()
    return elements


def build_with_python_docx(elements, output_path: pathlib.Path):
    doc = Document()
    for el in elements:
        if el["type"] == "heading":
            lvl = min(el.get("level", 1), 3)
            doc.add_heading(el["text"], level=lvl)
        elif el["type"] == "bullet":
            for item in el["items"]:
                p = doc.add_paragraph(item, style="List Bullet")
                p.paragraph_format.space_after = 3
        elif el["type"] == "table":
            header = el.get("header", [])
            rows = el.get("rows", [])
            table = doc.add_table(rows=len(rows) + 1, cols=max(len(header), len(rows[0]) if rows else 1))
            for idx, cell_val in enumerate(header):
                table.rows[0].cells[idx].text = cell_val
            for r_idx, row in enumerate(rows, start=1):
                for c_idx, cell_val in enumerate(row):
                    table.rows[r_idx].cells[c_idx].text = cell_val
        else:
            doc.add_paragraph(el["text"])
    doc.save(output_path)


# ---------- Fallback: build DOCX with stdlib ---------- #
NAMESPACES = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
for prefix, uri in NAMESPACES.items():
    ET.register_namespace(prefix, uri)


def _w(tag):
    return f"{{{NAMESPACES['w']}}}{tag}"


def paragraph_xml(text, style=None):
    p = ET.Element(_w("p"))
    if style:
        pPr = ET.SubElement(p, _w("pPr"))
        pStyle = ET.SubElement(pPr, _w("pStyle"))
        pStyle.set(_w("val"), style)
    r = ET.SubElement(p, _w("r"))
    t = ET.SubElement(r, _w("t"))
    t.text = text
    return p


def table_xml(header, rows):
    tbl = ET.Element(_w("tbl"))
    tblPr = ET.SubElement(tbl, _w("tblPr"))
    tblBorders = ET.SubElement(tblPr, _w("tblBorders"))
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        b = ET.SubElement(tblBorders, _w(side))
        b.set(_w("val"), "single")
        b.set(_w("sz"), "4")
        b.set(_w("space"), "0")
        b.set(_w("color"), "auto")

    tblGrid = ET.SubElement(tbl, _w("tblGrid"))
    col_count = max(len(header), max((len(r) for r in rows), default=0))
    for _ in range(col_count):
        ET.SubElement(tblGrid, _w("gridCol")).set(_w("w"), "4000")

    def add_row(cells, is_header=False):
        tr = ET.SubElement(tbl, _w("tr"))
        for val in cells:
            tc = ET.SubElement(tr, _w("tc"))
            p = paragraph_xml(val, style="Heading3" if is_header else None)
            tc.append(p)
            ET.SubElement(tc, _w("tcPr"))

    add_row(header, is_header=True)
    for r in rows:
        add_row(r, is_header=False)
    return tbl


def styles_xml():
    # Minimal styles for Normal, Heading1-3, ListParagraph
    styles = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="paragraph" w:styleId="Normal">
    <w:name w:val="Normal"/>
    <w:rsid w:val="00000000"/>
  </w:style>
  <w:style w:type="paragraph" w:styleId="Heading1">
    <w:name w:val="heading 1"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:uiPriority w:val="9"/><w:qFormat/>
    <w:pPr><w:keepNext/><w:keepLines/><w:outlineLvl w:val="0"/></w:pPr>
    <w:rPr><w:b/><w:sz w:val="32"/></w:rPr>
  </w:style>
  <w:style w:type="paragraph" w:styleId="Heading2">
    <w:name w:val="heading 2"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:uiPriority w:val="9"/><w:qFormat/>
    <w:pPr><w:keepNext/><w:keepLines/><w:outlineLvl w:val="1"/></w:pPr>
    <w:rPr><w:b/><w:sz w:val="28"/></w:rPr>
  </w:style>
  <w:style w:type="paragraph" w:styleId="Heading3">
    <w:name w:val="heading 3"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:uiPriority w:val="9"/><w:qFormat/>
    <w:pPr><w:keepNext/><w:keepLines/><w:outlineLvl w:val="2"/></w:pPr>
    <w:rPr><w:b/><w:sz w:val="24"/></w:rPr>
  </w:style>
  <w:style w:type="paragraph" w:styleId="ListParagraph">
    <w:name w:val="List Paragraph"/><w:basedOn w:val="Normal"/><w:uiPriority w:val="34"/><w:qFormat/>
  </w:style>
</w:styles>
"""
    return styles


def content_types_xml():
    return """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
</Types>
"""


def rels_xml():
    return """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>
"""


def doc_rels_xml():
    return """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>
"""


def build_document_xml(elements):
    body = ET.Element(_w("body"))
    for el in elements:
        if el["type"] == "heading":
            level = el.get("level", 1)
            style = f"Heading{min(level,3)}"
            body.append(paragraph_xml(el["text"], style=style))
        elif el["type"] == "bullet":
            for item in el["items"]:
                body.append(paragraph_xml(f"• {item}", style="ListParagraph"))
        elif el["type"] == "table":
            header = el.get("header", [])
            rows = el.get("rows", [])
            body.append(table_xml(header, rows))
        else:
            body.append(paragraph_xml(el["text"], style=None))
    sectPr = ET.SubElement(body, _w("sectPr"))
    ET.SubElement(sectPr, _w("pgSz"), {_w("w"): "12240", _w("h"): "15840"})
    ET.SubElement(sectPr, _w("pgMar"), {_w("top"): "720", _w("right"): "720", _w("bottom"): "720", _w("left"): "720"})

    doc = ET.Element(_w("document"))
    doc.append(body)
    return '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + ET.tostring(doc, encoding="unicode")


def build_with_stdlib(elements, output_path: pathlib.Path):
    document_xml = build_document_xml(elements)
    with zipfile.ZipFile(output_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types_xml())
        zf.writestr("_rels/.rels", rels_xml())
        zf.writestr("word/document.xml", document_xml)
        zf.writestr("word/styles.xml", styles_xml())
        zf.writestr("word/_rels/document.xml.rels", doc_rels_xml())


def generate(md_path: pathlib.Path, out_path: pathlib.Path):
    text = md_path.read_text(encoding="utf-8")
    elements = parse_markdown(text)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    if Document:
        build_with_python_docx(elements, out_path)
    else:
        build_with_stdlib(elements, out_path)


def main():
    if len(sys.argv) < 3:
        print("Usage: py scripts/generate_docx.py <input.md> <output.docx>")
        sys.exit(1)
    md_path = pathlib.Path(sys.argv[1])
    out_path = pathlib.Path(sys.argv[2])
    if not md_path.exists():
        print(f"Input markdown not found: {md_path}")
        sys.exit(1)
    generate(md_path, out_path)
    print(f"Generated {out_path}")


if __name__ == "__main__":
    main()
